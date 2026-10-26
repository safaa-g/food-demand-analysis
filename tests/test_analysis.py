import sys
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

import food_demand_analysis as project
import create_presentation as presentation


class PriceFeatureTests(unittest.TestCase):
    def test_discount_features_keep_premium_prices_signed(self) -> None:
        source = pd.DataFrame(
            {"base_price": [100.0, 100.0, 0.0], "checkout_price": [80.0, 110.0, 25.0]}
        )
        result = project.calculate_price_features(source)
        self.assertEqual(result["discount_amount"].tolist()[:2], [20.0, -10.0])
        self.assertEqual(result["discount_rate"].iloc[0], 0.2)
        self.assertEqual(result["discount_rate"].iloc[1], -0.1)
        self.assertTrue(pd.isna(result["discount_rate"].iloc[2]))
        self.assertNotIn("discount_rate", source.columns)


class ContributorTests(unittest.TestCase):
    def test_top_contributors_returns_smallest_prefix_reaching_threshold(self) -> None:
        summary = pd.DataFrame({"total_orders": [50, 30, 20]})
        result = project.top_contributors(summary, threshold_pct=80)
        self.assertEqual(result["total_orders"].tolist(), [50, 30])

    def test_invalid_contribution_threshold_is_rejected(self) -> None:
        with self.assertRaises(ValueError):
            project.top_contributors(pd.DataFrame({"total_orders": [1]}), 0)


class SourceDataTests(unittest.TestCase):
    def test_source_files_load_and_validate(self) -> None:
        train, test, quality = project.load_and_validate(ROOT / "data" / "raw")
        self.assertEqual(quality["train_rows"], 456548)
        self.assertEqual(quality["test_rows"], 32573)
        self.assertEqual(quality["train_week_range"], [1, 145])
        self.assertEqual(quality["test_week_range"], [146, 155])
        self.assertIn("category", train.columns)
        self.assertIn("cuisine", test.columns)


class ForecastingTests(unittest.TestCase):
    def test_models_use_time_holdout_and_write_nonnegative_predictions(self) -> None:
        train = pd.DataFrame(
            {
                "id": range(290),
                "week": [week for week in range(1, 146) for _ in range(2)],
                "checkout_price": [100.0, 90.0] * 145,
                "base_price": [110.0, 100.0] * 145,
                "discount_amount": [10.0] * 290,
                "discount_rate": [10 / 110, 0.1] * 145,
                "emailer_for_promotion": [0, 1] * 145,
                "homepage_featured": [1, 0] * 145,
                "center_id": [13, 11] * 145,
                "city_code": [590, 679] * 145,
                "region_code": [56, 56] * 145,
                "center_type": ["TYPE_B", "TYPE_A"] * 145,
                "op_area": [6.7, 3.7] * 145,
                "meal_id": [1885, 1993] * 145,
                "category": ["Beverages", "Beverages"] * 145,
                "cuisine": ["Thai", "Thai"] * 145,
                "num_orders": [100 + week for week in range(1, 146) for _ in range(2)],
            }
        )
        test = train.iloc[:2].drop(columns="num_orders").copy()
        test["id"] = [1001, 1002]
        test["week"] = [146, 147]

        with tempfile.TemporaryDirectory() as temp_dir:
            scores = project.train_and_predict(train, test, Path(temp_dir))
            submission = pd.read_csv(Path(temp_dir) / "submission.csv")

        self.assertEqual(set(scores["model"]), {"LinearRegression", "RandomForestRegressor"})
        self.assertEqual(submission["id"].tolist(), [1001, 1002])
        self.assertTrue((submission["num_orders"] >= 0).all())


class PresentationTests(unittest.TestCase):
    def test_powerpoint_has_expected_slides_images_and_fits_canvas(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            output = presentation.build_presentation(Path(temp_dir) / "deck.pptx")
            deck = Presentation(output)

        self.assertEqual(len(deck.slides), 7)
        embedded_images = 0
        for slide in deck.slides:
            for shape in slide.shapes:
                self.assertGreaterEqual(shape.left, 0)
                self.assertGreaterEqual(shape.top, 0)
                self.assertLessEqual(shape.left + shape.width, deck.slide_width)
                self.assertLessEqual(shape.top + shape.height, deck.slide_height)
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    embedded_images += 1
        self.assertEqual(embedded_images, 4)


if __name__ == "__main__":
    unittest.main()
